"""Tests for auto-ppt-playbook-builder/ppt_core.py. No Streamlit needed.

Run:  python3 tests/run_tests.py   (from the app directory)
"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PIL import Image  # noqa: E402
from pptx import Presentation  # noqa: E402

import ppt_core as core  # noqa: E402

passed = 0


def check(name, cond, detail=""):
    global passed
    if not cond:
        print(f"FAIL: {name} {detail}")
        sys.exit(1)
    passed += 1
    print(f"ok: {name}")


def reopen(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


# --- annotation -------------------------------------------------------
img = Image.new("RGB", (800, 600), "white")
ann = core.annotate_screenshot(img, [{"x": 50, "y": 50, "n": 1}])
px = ann.getpixel((412, 300))  # inside badge, off the "1" glyph
check("badge drawn red near center",
      px[0] > 150 and px[1] < 100, str(px))
check("corner untouched", ann.getpixel((5, 5)) == (255, 255, 255))
check("input image not mutated", img.getpixel((400, 300)) == (255, 255, 255))
check("no callouts -> identical copy",
      list(core.annotate_screenshot(img, []).getdata()) == list(img.getdata()))

# --- annotated slide --------------------------------------------------
prs = core.new_presentation()
core.add_annotated_slide(prs, ann, [{"n": 1, "text": "Click here"}],
                         title="Walkthrough")
data = core.prs_to_bytes(prs)
back = reopen(data)
check("one annotated slide", len(back.slides) == 1)
shapes = back.slides[0].shapes
check("slide has picture + table",
      any(s.shape_type == 13 for s in shapes)  # 13 = picture
      and any(s.has_table for s in shapes))
tbl = next(s.table for s in shapes if s.has_table)
check("steps table shape", len(tbl.rows) == 2 and len(tbl.columns) == 2)
check("steps table content", tbl.cell(1, 1).text == "Click here")

# --- dictionary pagination --------------------------------------------
prs = core.new_presentation()
rows = [[f"M{i}", f"SUM(X{i})", f"desc {i}"] for i in range(15)]
n = core.add_dictionary_slides(prs, ["Measure", "DAX", "Description"], rows,
                               rows_per_slide=7, title_prefix="DD")
check("pagination count", n == 3)
back = reopen(core.prs_to_bytes(prs))
check("3 slides created", len(back.slides) == 3)
t1 = next(s.table for s in back.slides[0].shapes if s.has_table)
t3 = next(s.table for s in back.slides[2].shapes if s.has_table)
check("page 1 has 7 data rows", len(t1.rows) == 8)
check("page 3 has 1 data row", len(t3.rows) == 2)
check("last row content", t3.cell(1, 0).text == "M14")

# --- mail merge --------------------------------------------------------
tpl = core.new_presentation()
slide = tpl.slides.add_slide(core.blank_layout(tpl))
tx = slide.shapes.add_textbox(core.Inches(1), core.Inches(1),
                              core.Inches(5), core.Inches(1))
tx.text_frame.text = "Store: {{Store}} — Manager {{ Manager }}"
tshape = slide.shapes.add_table(2, 2, core.Inches(1), core.Inches(3),
                                core.Inches(5), core.Inches(1))
tshape.table.cell(0, 0).text = "{{Store}}"
tshape.table.cell(0, 1).text = "static"
tshape.table.cell(1, 0).text = "{{Missing}}"
tpl_bytes = core.prs_to_bytes(tpl)

tags = core.find_tags(core.load_template(tpl_bytes), 0)
check("tag detection", tags == ["Store", "Manager", "Missing"], str(tags))

# template with an image (red square) to verify rel retargeting
img_buf = io.BytesIO()
Image.new("RGB", (50, 50), "red").save(img_buf, format="PNG")
img_buf.seek(0)
slide.shapes.add_picture(img_buf, core.Inches(7), core.Inches(1),
                         width=core.Inches(1), height=core.Inches(1))
tpl_bytes = core.prs_to_bytes(tpl)

out = core.mail_merge(tpl_bytes, 0, [{"Store": "A", "Manager": "Amy"},
                                     {"Store": "B", "Manager": "Bob"}])
merged = reopen(out)
check("merged slide count", len(merged.slides) == 2)
texts = [s.text_frame.text for sh in merged.slides[0].shapes
         for s in [sh] if sh.has_text_frame]
check("tags replaced slide 1", any("Store: A — Manager Amy" in t for t in texts))
tbl0 = next(s.table for s in merged.slides[0].shapes if s.has_table)
check("table tag replaced", tbl0.cell(0, 0).text == "A")
check("missing tag left as-is", tbl0.cell(1, 0).text == "{{Missing}}")
check("image preserved on clone",
      sum(1 for s in merged.slides[0].shapes if s.shape_type == 13) == 1)
texts1 = [sh.text_frame.text for sh in merged.slides[1].shapes
          if sh.has_text_frame]
check("tags replaced slide 2", any("Store: B" in t for t in texts1))

# --- markdown + csv -----------------------------------------------------
cols, rows = core.parse_markdown_table(
    "| Measure | DAX |\n|---|---|\n| Total | SUM(x) | extra |\n| Avg | AVG(x) |")
check("markdown parse", cols == ["Measure", "DAX"] and rows[0] == ["Total", "SUM(x)"])
try:
    core.parse_markdown_table("no table here")
    check("markdown rejects non-table", False)
except ValueError:
    check("markdown rejects non-table", True)

df = core.read_csv_smart("a,b\njosé,1\n".encode("cp1252"))
check("csv cp1252 fallback", df.iloc[0, 0] == "josé")

print(f"\nAll {passed} checks passed.")
