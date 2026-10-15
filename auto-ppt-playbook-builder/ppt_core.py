"""Core logic for the PPT Playbook Builder. Pure — no ``st.*`` calls.

Covers: screenshot annotation (Pillow), annotated-slide building,
paginated data-dictionary table slides, and template mail-merge
(slide cloning + ``{{tag}}`` replacement), all on python-pptx.
"""

import copy
import io
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

TAG_RE = re.compile(r"\{\{\s*([A-Za-z0-9_ ]+?)\s*\}\}")

BADGE_FILL = (211, 47, 47)      # corporate red
HEADER_FILL = RGBColor(0x1F, 0x49, 0x7D)   # deep navy
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x33, 0x33, 0x33)
LIGHT_ROW = RGBColor(0xF2, 0xF2, 0xF2)


# ------------------------------------------------------------- fonts
def _load_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                continue
    return ImageFont.load_default()


# ------------------------------------------------------- annotation
def annotate_screenshot(img: Image.Image,
                        callouts: list[dict]) -> Image.Image:
    """Draw numbered badges on a screenshot.

    callouts: [{"x": 0-100, "y": 0-100, "n": 1}, ...]
    Returns a new RGB image; the input is not modified.
    """
    out = img.convert("RGB")
    w, h = out.size
    if not callouts:
        return out
    radius = max(16, int(min(w, h) * 0.032))
    font = _load_font(int(radius * 1.1))
    draw = ImageDraw.Draw(out)
    for c in callouts:
        cx = int(w * min(100, max(0, c["x"])) / 100)
        cy = int(h * min(100, max(0, c["y"])) / 100)
        n = str(c.get("n", ""))
        # white halo + red badge for contrast on any background
        draw.ellipse([cx - radius - 3, cy - radius - 3,
                      cx + radius + 3, cy + radius + 3], fill="white")
        draw.ellipse([cx - radius, cy - radius,
                      cx + radius, cy + radius], fill=BADGE_FILL)
        bbox = draw.textbbox((0, 0), n, font=font)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.text((cx - tw / 2 - bbox[0], cy - th / 2 - bbox[1]),
                  n, font=font, fill="white")
    return out


# ------------------------------------------------------- presentation
def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    return prs


def load_template(template_bytes: bytes | None) -> Presentation:
    if template_bytes:
        return Presentation(io.BytesIO(template_bytes))
    return new_presentation()


def blank_layout(prs: Presentation):
    """The emptiest layout available (no placeholders if possible)."""
    for layout in prs.slide_layouts:
        if not layout.placeholders:
            return layout
    if len(prs.slide_layouts) > 6:
        return prs.slide_layouts[6]
    return prs.slide_layouts[0]


def prs_to_bytes(prs: Presentation) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _add_title_box(slide, text: str, top=Inches(0.15)):
    txBox = slide.shapes.add_textbox(Inches(0.4), top, Inches(12.5), Inches(0.6))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = HEADER_FILL
    return txBox


def _style_header_cell(cell, text: str, size: int = 11):
    cell.fill.solid()
    cell.fill.fore_color.rgb = HEADER_FILL
    cell.text_frame.word_wrap = True
    p = cell.text_frame.paragraphs[0]
    p.text = str(text)
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = WHITE


def _style_body_cell(cell, text: str, size: int = 10, shade: bool = False):
    if shade:
        cell.fill.solid()
        cell.fill.fore_color.rgb = LIGHT_ROW
    cell.text_frame.word_wrap = True
    p = cell.text_frame.paragraphs[0]
    p.text = str(text)
    p.font.size = Pt(size)
    p.font.color.rgb = DARK


def add_annotated_slide(prs: Presentation,
                        annotated_img: Image.Image,
                        steps: list[dict],
                        title: str = "Annotated Walkthrough") -> None:
    """Left: annotated screenshot. Right: numbered explanation table."""
    slide = prs.slides.add_slide(blank_layout(prs))
    _add_title_box(slide, title)

    # Screenshot (left ~58%)
    buf = io.BytesIO()
    annotated_img.save(buf, format="PNG")
    buf.seek(0)
    left, top = Inches(0.4), Inches(1.0)
    pic_w = Inches(7.3)
    iw, ih = annotated_img.size
    pic_h = pic_w * ih / iw
    max_h = Inches(6.0)
    if pic_h > max_h:
        pic_h = max_h
        pic_w = pic_h * iw / ih
    slide.shapes.add_picture(buf, left, top, width=pic_w, height=pic_h)

    # Steps table (right)
    rows = max(1, len(steps))
    tbl_left, tbl_top = Inches(8.1), Inches(1.0)
    tbl_w = Inches(4.8)
    shape = slide.shapes.add_table(rows + 1, 2, tbl_left, tbl_top, tbl_w,
                                   Inches(0.45 * (rows + 1)))
    table = shape.table
    table.columns[0].width = Inches(0.6)
    table.columns[1].width = Inches(4.2)
    _style_header_cell(table.cell(0, 0), "#")
    _style_header_cell(table.cell(0, 1), "Step")
    for i, s in enumerate(steps, start=1):
        _style_body_cell(table.cell(i, 0), str(s.get("n", i)),
                         size=12, shade=i % 2 == 0)
        p = table.cell(i, 0).text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.font.bold = True
        _style_body_cell(table.cell(i, 1), s.get("text", ""), shade=i % 2 == 0)
    if not steps:
        _style_body_cell(table.cell(1, 0), "—")
        _style_body_cell(table.cell(1, 1), "No steps defined.")


def add_dictionary_slides(prs: Presentation,
                          columns: list[str],
                          rows: list[list],
                          rows_per_slide: int = 7,
                          title_prefix: str = "Data Dictionary") -> int:
    """Paginate rows into table slides. Returns the number of slides added."""
    if not columns:
        raise ValueError("No columns to render.")
    chunks = [rows[i:i + rows_per_slide]
              for i in range(0, len(rows), rows_per_slide)] or [[]]
    total = len(chunks)
    size = 12 if len(columns) <= 3 else 10
    for idx, chunk in enumerate(chunks, start=1):
        slide = prs.slides.add_slide(blank_layout(prs))
        _add_title_box(slide, f"{title_prefix} ({idx}/{total})")
        n_rows = max(1, len(chunk))
        shape = slide.shapes.add_table(
            n_rows + 1, len(columns), Inches(0.4), Inches(1.0),
            Inches(12.5), Inches(0.4 * (n_rows + 1)))
        table = shape.table
        for j, col in enumerate(columns):
            _style_header_cell(table.cell(0, j), col, size=size)
        for i, row in enumerate(chunk, start=1):
            for j in range(len(columns)):
                val = row[j] if j < len(row) else ""
                _style_body_cell(table.cell(i, j), val, size=size,
                                 shade=i % 2 == 0)
        # Even column widths
        for j in range(len(columns)):
            table.columns[j].width = Inches(12.5 / len(columns))
    return total


# ------------------------------------------------------- mail merge
def _iter_text_frames(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def find_tags(prs: Presentation, slide_idx: int = 0) -> list[str]:
    """All {{tag}} names found on one template slide, in first-seen order."""
    tags: list[str] = []
    slide = prs.slides[slide_idx]
    for tf in _iter_text_frames(slide):
        for para in tf.paragraphs:
            full = "".join(r.text for r in para.runs)
            for m in TAG_RE.finditer(full):
                name = m.group(1).strip()
                if name not in tags:
                    tags.append(name)
    return tags


def _replace_in_text_frame(tf, mapping: dict) -> None:
    for para in tf.paragraphs:
        full = "".join(r.text for r in para.runs)
        if "{{" not in full:
            continue
        new_full = full
        for key, val in mapping.items():
            new_full = new_full.replace("{{" + key + "}}", str(val))
            new_full = new_full.replace("{{ " + key + " }}", str(val))
        if new_full != full and para.runs:
            para.runs[0].text = new_full
            for r in para.runs[1:]:
                r.text = ""


_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"


def clone_slide(prs: Presentation, src_slide) -> object:
    """Deep-copy a slide (shapes + image relationships)."""
    dest = prs.slides.add_slide(blank_layout(prs))
    # Blank layouts can carry placeholders; strip them for a clean clone.
    for shp in list(dest.shapes):
        el = shp._element
        el.getparent().remove(el)
    src_part, dest_part = src_slide.part, dest.part
    for shape in src_slide.shapes:
        el = copy.deepcopy(shape._element)
        for blip in el.iterfind(f".//{_BLIP}"):
            rid = blip.get(_EMBED)
            if rid and rid in src_part.rels:
                new_rid = dest_part.relate_to(src_part.rels[rid].target_part,
                                              RT.IMAGE)
                blip.set(_EMBED, new_rid)
        dest.shapes._spTree.insert_element_before(el, "p:extLst")
    return dest


def mail_merge(template_bytes: bytes,
               template_slide_idx: int,
               dataset: list[dict]) -> bytes:
    """One cloned slide per dataset row, with {{tags}} replaced.

    dataset: list of {tag: value}. Missing tags are left as-is.
    """
    if not dataset:
        raise ValueError("Dataset is empty — nothing to merge.")
    prs = Presentation(io.BytesIO(template_bytes))
    if not (0 <= template_slide_idx < len(prs.slides)):
        raise ValueError(f"Template slide #{template_slide_idx + 1} does not "
                         f"exist (template has {len(prs.slides)} slide(s)).")
    src = prs.slides[template_slide_idx]
    for row in dataset:
        mapping = {str(k).strip(): ("" if v is None else v)
                   for k, v in row.items()}
        new_slide = clone_slide(prs, src)
        for tf in _iter_text_frames(new_slide):
            _replace_in_text_frame(tf, mapping)
    # Remove the original template slide(s); keep only merged output
    # (delete in reverse so indices stay valid).
    for idx in sorted(range(len(prs.slides) - len(dataset)), reverse=True):
        prs.part.drop_rel(prs.slides._sldIdLst[idx].rId)
        del prs.slides._sldIdLst[idx]
    return prs_to_bytes(prs)


# ------------------------------------------------------- csv helper
def read_csv_smart(raw: bytes):
    """Encoding-tolerant CSV reader (utf-8-sig -> utf-8 -> cp1252)."""
    import pandas as pd
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return pd.read_csv(io.BytesIO(raw), dtype=str, encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return pd.read_csv(io.BytesIO(raw), dtype=str, encoding="latin-1")


# ------------------------------------------------------- markdown table
def parse_markdown_table(text: str) -> tuple[list[str], list[list[str]]]:
    """Parse a GitHub-style markdown table into (columns, rows)."""
    lines = [l.strip() for l in text.strip().splitlines() if l.strip()]
    rows: list[list[str]] = []
    for line in lines:
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(set(c) <= set(":- ") for c in cells):
            continue  # separator row
        rows.append(cells)
    if not rows:
        raise ValueError("No markdown table rows found (need | cell | rows).")
    columns = rows[0]
    data = rows[1:]
    width = len(columns)
    data = [r + [""] * (width - len(r)) if len(r) < width else r[:width]
            for r in data]
    return columns, data
