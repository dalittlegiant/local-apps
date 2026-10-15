"""Core logic for the Messy Excel Sanitizer. Pure — no ``st.*`` calls.

Pipeline: read (with merged-cell handling) -> skip title rows -> header
row -> sanitization (whitespace, header case, date unification, empties,
duplicates) -> schema validation -> concatenation.
"""

import io
import re

import pandas as pd
from pandas.api.types import is_string_dtype

DATE_LIKE = re.compile(r"\d{1,4}[-/\.]\d{1,2}[-/\.]\d{1,4}")


def _is_text(s: pd.Series) -> bool:
    # object, str, and StringDtype (pandas 3 infers str by default)
    return is_string_dtype(s)


# ------------------------------------------------------------ reading
def _read_csv_raw(raw: bytes) -> pd.DataFrame:
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return pd.read_csv(io.BytesIO(raw), header=None, dtype=object,
                               encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return pd.read_csv(io.BytesIO(raw), header=None, dtype=object,
                       encoding="latin-1")


def read_source(raw: bytes, filename: str, sheet: str | None = None,
                skip_rows: int = 0, unmerge: bool = False) -> pd.DataFrame:
    """Read one (file, sheet) source into a header-less DataFrame.

    xlsx with unmerge=True: merged ranges are unmerged via openpyxl first
    (values are NOT filled — the pipeline forward-fills afterwards).
    Raises ValueError on failure.
    """
    lname = filename.lower()
    try:
        if lname.endswith((".xlsx", ".xls")):
            data = raw
            if unmerge:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
                ws = wb[sheet] if sheet in wb.sheetnames else wb.active
                for mr in list(ws.merged_cells.ranges):
                    ws.unmerge_cells(str(mr))
                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                data = buf.getvalue()
            df = pd.read_excel(io.BytesIO(data), sheet_name=sheet or 0,
                               header=None, dtype=object)
        elif lname.endswith(".csv"):
            df = _read_csv_raw(raw)
        else:
            raise ValueError(f"Unsupported file type: {filename}")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not read {filename}: {exc}") from exc

    if skip_rows:
        df = df.iloc[skip_rows:].reset_index(drop=True)
    if df.empty:
        raise ValueError(f"{filename}: no rows left after skipping "
                         f"{skip_rows} title row(s).")
    header = [(str(v).strip() if pd.notna(v) else "") for v in df.iloc[0]]
    if unmerge:
        # Horizontally merged headers leave gaps — fill across the header row.
        header = (pd.Series(header).replace("", pd.NA).ffill()
                  .fillna("").tolist())
    df = df.iloc[1:].reset_index(drop=True)
    df.columns = _dedupe([h if h else f"col_{i}" for i, h in enumerate(header)])
    return df


def _dedupe(cols: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for c in cols:
        if c in seen:
            seen[c] += 1
            out.append(f"{c}_{seen[c]}")
        else:
            seen[c] = 1
            out.append(c)
    return out


# -------------------------------------------------------- sanitizing
def _snake(name: str) -> str:
    name = name.strip().lower()
    name = re.sub(r"[^0-9a-z]+", "_", name).strip("_")
    return name or "col"


def sanitize_df(df: pd.DataFrame, opts: dict) -> tuple[pd.DataFrame, list[str]]:
    """Apply the sanitization pipeline. Returns (clean_df, notes)."""
    notes: list[str] = []
    df = df.copy()

    if opts.get("ffill"):
        df = df.ffill()
        notes.append("Forward-filled values down (unmerged cells).")
    if opts.get("strip_ws", True):
        for c in df.columns:
            if _is_text(df[c]):
                df[c] = df[c].apply(
                    lambda v: v.strip() if isinstance(v, str) else v)
        notes.append("Stripped leading/trailing whitespace in text columns.")
    if opts.get("header_mode") == "snake":
        df.columns = _dedupe([_snake(str(c)) for c in df.columns])
        notes.append("Headers standardized to snake_case.")
    elif opts.get("header_mode") == "strip":
        df.columns = _dedupe([str(c).strip() for c in df.columns])

    if opts.get("date_unify"):
        dayfirst = bool(opts.get("dayfirst"))
        converted = []
        for c in df.columns:
            s = df[c]
            if not _is_text(s):
                continue
            sample = s.dropna().astype(str)
            if sample.empty:
                continue
            datelike = sample.str.contains(DATE_LIKE, regex=True).mean()
            if datelike >= 0.5:
                parsed = pd.to_datetime(s, errors="coerce", dayfirst=dayfirst)
                if parsed.notna().sum() >= 0.5 * s.notna().sum():
                    df[c] = parsed
                    converted.append(str(c))
        if converted:
            notes.append("Unified date formats in: " + ", ".join(converted))

    if opts.get("drop_empty_cols", True):
        before = len(df.columns)
        df = df.dropna(axis=1, how="all")
        if len(df.columns) < before:
            notes.append(f"Dropped {before - len(df.columns)} fully-empty column(s).")
    if opts.get("drop_empty_rows", True):
        before = len(df)
        df = df.dropna(axis=0, how="all").reset_index(drop=True)
        if len(df) < before:
            notes.append(f"Dropped {before - len(df)} fully-empty row(s).")
    if opts.get("drop_dupes"):
        before = len(df)
        df = df.drop_duplicates().reset_index(drop=True)
        if len(df) < before:
            notes.append(f"Dropped {before - len(df)} duplicate row(s).")
    return df, notes


# ------------------------------------------------- validation & merge
def validate_schema(df: pd.DataFrame,
                    expected: list[str]) -> tuple[list[str], list[str]]:
    """Returns (missing_columns, extra_columns). Case/space-insensitive."""
    norm = lambda cols: {re.sub(r"\s+", "", str(c)).lower(): str(c) for c in cols}
    have, want = norm(df.columns), norm(expected)
    missing = [want[k] for k in want if k not in have]
    extra = [have[k] for k in have if k not in want]
    return missing, extra


def concat_sources(frames: list[tuple[str, str, pd.DataFrame]]) -> pd.DataFrame:
    """Concatenate (file, sheet, df) with union of columns + source columns."""
    parts = []
    for fname, sheet, df in frames:
        d = df.copy()
        d["_source_file"] = fname
        d["_source_sheet"] = sheet or ""
        parts.append(d)
    if not parts:
        raise ValueError("No sources to concatenate.")
    return pd.concat(parts, ignore_index=True, sort=False)
